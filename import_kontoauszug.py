"""CLI: Kontoauszug-PDF -> regelbasierte Kategorisierung -> Kontrolle ->
Anhaengen an die "Transaktionen"-Tabelle in Finanzen.xlsx. Kein API-Key,
kein Server, keine Kosten.

Unterstuetzt automatisch alle drei kalibrierten PDF-Formate (Girokonto-
Umsaetze, Kreditkartenkonto-Umsaetze, klassische Kreditkarten-Umsatzaufstellung).

Nutzung:
    python import_kontoauszug.py <pfad_zur_pdf> [--quelle Girokonto|Kreditkarte]

Beispiel:
    python import_kontoauszug.py Umsaetze_Oktober.pdf
"""
from __future__ import annotations

import argparse
import sys

import openpyxl

import categorizer_rules
import pdf_extractor
import pdf_transaction_parser
import xlsx_writer

DEFAULT_XLSX = r"C:\Users\Nils\Downloads\Finanzen.xlsx"


def existing_keys(xlsx_path: str) -> set[tuple]:
    """(datum, betrag, verwendungszweck) bereits vorhandener Zeilen fuer
    einfache Duplikat-Erkennung."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[xlsx_writer.SHEET_NAME]
    keys = set()
    for row in range(2, ws.max_row + 1):
        datum, _monat, _typ, _kat, betrag, zweck = (ws.cell(row=row, column=c).value for c in range(1, 7))
        if datum is not None:
            keys.add((str(datum), betrag, str(zweck or "")))
    return keys


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pdf_path")
    parser.add_argument("--xlsx", default=DEFAULT_XLSX)
    parser.add_argument("--quelle", choices=["Girokonto", "Kreditkarte"], default=None,
                         help="Wird bei Girokonto/Kreditkartenkonto-Format automatisch erkannt; "
                              "nur bei der alten Einzel-Kreditkarten-Umsatzaufstellung noetig.")
    parser.add_argument("--yes", action="store_true", help="ohne Rueckfrage eintragen")
    args = parser.parse_args()

    raw_text = pdf_extractor.extract_text(args.pdf_path)
    parsed = pdf_transaction_parser.parse_transactions(raw_text)

    if not parsed:
        print("Keine Buchungszeilen erkannt. Regex in pdf_transaction_parser.py "
              "muss vermutlich an dieses Format angepasst werden.")
        sys.exit(1)

    quelle_default = args.quelle or ("Kreditkarte" if "Kreditkartenkonto" in raw_text else "Girokonto")
    for tx in parsed:
        is_income = tx["betrag"] > 0
        match_text = tx.get("match_text", tx["verwendungszweck"])
        tx["kategorie"] = categorizer_rules.categorize(match_text, is_income)
        tx.setdefault("quelle", quelle_default)

    dupes = existing_keys(args.xlsx)

    print(f"\n{len(parsed)} Buchungen erkannt:\n")
    print(f"{'Datum':<12} {'Betrag':>10}  {'Kategorie':<15} {'Verwendungszweck':<35} Hinweis")
    print("-" * 100)
    for tx in parsed:
        key = (tx["datum"], tx["betrag"], tx["verwendungszweck"])
        hinweise = []
        if key in dupes:
            hinweise.append("DUPLIKAT?")
        if tx.get("unsicher"):
            hinweise.append("Vorzeichen unsicher")
        print(f"{tx['datum']:<12} {tx['betrag']:>10.2f}  {tx['kategorie']:<15} "
              f"{tx['verwendungszweck'][:35]:<35} {', '.join(hinweise)}")

    if not args.yes:
        antwort = input("\nDiese Zeilen an die Transaktionen-Tabelle anhaengen? [j/N] ").strip().lower()
        if antwort != "j":
            print("Abgebrochen, nichts geschrieben.")
            return

    result = xlsx_writer.write_transactions(args.xlsx, parsed)
    print(f"\nFertig: {result['geschrieben']} Buchungen angehaengt.")


if __name__ == "__main__":
    main()
