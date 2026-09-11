"""Schreibt bestaetigte Transaktionen direkt in die Finanzen.xlsx (kein
Google Sheets, kein Server) -- an die exakte Tabellenstruktur angepasst,
wie sie in der Datei vorgefunden wurde (Header in Zeile 40, Daten ab
Zeile 42, Ausgaben in Spalte A-D, Einnahmen in Spalte F-I)."""
from __future__ import annotations

from datetime import date, datetime
from typing import Any

import openpyxl

DATE_NUMBER_FORMAT = "dd.mm"
DATA_START_ROW = 42
# SUMIFS-Formeln in den Monatsblaettern decken diese Bereiche ab --
# darueber hinaus wuerden neue Zeilen nicht mehr mitgezaehlt.
AUSGABEN_MAX_ROW = 133
EINNAHMEN_MAX_ROW = 52

MONTH_SHEET_NAMES = {
    "januar": "Januar ", "februar": "Februar ", "märz": "März ", "maerz": "März ",
    "april": "April ", "mai": "Mai ", "juni": "Juni ", "juli": "Juli ",
    "august": "August ", "september": "September ", "oktober": "Oktober ",
    "november": "November ", "dezember": "Dezember ",
}


def resolve_sheet_name(month: str) -> str:
    key = month.strip().lower()
    if key not in MONTH_SHEET_NAMES:
        raise ValueError(f"Unbekannter Monat: {month!r}")
    return MONTH_SHEET_NAMES[key]


def _next_empty_row(ws, col: str, start_row: int, max_row: int) -> int:
    row = start_row
    while row <= max_row and ws[f"{col}{row}"].value is not None:
        row += 1
    if row > max_row:
        raise RuntimeError(
            f"Kein Platz mehr in Spalte {col} (Zeile {start_row}-{max_row} voll). "
            "SUMIFS-Bereich im Sheet muesste erweitert werden."
        )
    return row


def _as_date(value: Any) -> date:
    if isinstance(value, (date, datetime)):
        return value
    return datetime.strptime(str(value), "%Y-%m-%d").date()


def write_transactions(
    xlsx_path: str,
    month: str,
    ausgaben: list[dict[str, Any]],
    einnahmen: list[dict[str, Any]],
) -> dict[str, int]:
    """ausgaben/einnahmen: Liste von Dicts mit Keys datum, betrag, kategorie,
    verwendungszweck. Schreibt fortlaufend ab der ersten freien Zeile."""
    sheet_name = resolve_sheet_name(month)
    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    ws = wb[sheet_name]

    row = _next_empty_row(ws, "A", DATA_START_ROW, AUSGABEN_MAX_ROW)
    for tx in ausgaben:
        ws[f"A{row}"] = _as_date(tx["datum"])
        ws[f"A{row}"].number_format = DATE_NUMBER_FORMAT
        ws[f"B{row}"] = tx["betrag"]
        ws[f"C{row}"] = tx["kategorie"]
        ws[f"D{row}"] = tx.get("verwendungszweck", "")
        row += 1

    row = _next_empty_row(ws, "F", DATA_START_ROW, EINNAHMEN_MAX_ROW)
    for tx in einnahmen:
        ws[f"F{row}"] = _as_date(tx["datum"])
        ws[f"F{row}"].number_format = DATE_NUMBER_FORMAT
        ws[f"G{row}"] = tx["betrag"]
        ws[f"H{row}"] = tx["kategorie"]
        ws[f"I{row}"] = tx.get("verwendungszweck", "")
        row += 1

    wb.save(xlsx_path)
    return {"ausgaben_geschrieben": len(ausgaben), "einnahmen_geschrieben": len(einnahmen)}
